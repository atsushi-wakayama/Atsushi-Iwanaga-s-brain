# -*- coding: utf-8 -*-
"""gobo_intersections.js から、警察に緯度経度を依頼するための交差点一覧(xlsx/csv)を生成する。"""
import json, os, subprocess
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

OUT_DIR = '/Users/iwanaga/claude code用/交通事故マップ'
JS = os.path.join(OUT_DIR, 'gobo_intersections.js')

# JS の配列を node 経由で JSON 取得
data = json.loads(subprocess.check_output(
    ['node', '-e', f'global.window={{}};require({json.dumps(JS)});process.stdout.write(JSON.stringify(window.OFFICIAL_HOTSPOTS))']
).decode())

YEARS = [2021, 2022, 2023, 2024, 2025]
WAREKI = {2021:'R3',2022:'R4',2023:'R5',2024:'R6',2025:'R7'}

# 出現年数の多い順→最多件数の多い順
def maxcount(d): return max(v['count'] for v in d['years'].values())
data.sort(key=lambda d:(-len(d['years']), -maxcount(d)))

headers = ['No','交差点名','市町村','主路線']
for y in YEARS:
    headers.append(f'{y}({WAREKI[y]})\n順位/件数')
headers += ['出現年数','最多件数','推定緯度(参考)','推定経度(参考)','確度',
            '警察提供 緯度','警察提供 経度','備考']

rows = []
for i, d in enumerate(data, 1):
    row = [i, d['name'], d['city'], d['route']]
    for y in YEARS:
        v = d['years'].get(str(y)) or d['years'].get(y)
        row.append(f"{v['rank']}位/{v['count']}件" if v else '')
    row += [len(d['years']), maxcount(d),
            d['lat'], d['lon'], ('確度高' if d['conf']=='ok' else '要確認'),
            '', '', '']
    rows.append(row)

# --- xlsx ---
wb = Workbook(); ws = wb.active; ws.title = '御坊署 多発交差点一覧'
ws.append(headers)
for r in rows: ws.append(r)

# 体裁
hdr_fill = PatternFill('solid', fgColor='1D3557')
req_fill = PatternFill('solid', fgColor='FFF3CD')  # 警察記入欄を強調
thin = Side(style='thin', color='CCCCCC')
border = Border(left=thin, right=thin, top=thin, bottom=thin)
for c, h in enumerate(headers, 1):
    cell = ws.cell(1, c)
    cell.font = Font(bold=True, color='FFFFFF', size=10)
    cell.fill = hdr_fill
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
req_cols = [headers.index('警察提供 緯度')+1, headers.index('警察提供 経度')+1]
for r in range(2, len(rows)+2):
    for c in range(1, len(headers)+1):
        cell = ws.cell(r, c); cell.border = border
        cell.alignment = Alignment(horizontal='center', vertical='center')
        if c in req_cols: cell.fill = req_fill
widths = [4,16,8,16, 11,11,11,11,11, 8,8,13,13,7, 13,13,16]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[ws.cell(1,i).column_letter].width = w
ws.freeze_panes = 'B2'; ws.row_dimensions[1].height = 30

xlsx = os.path.join(OUT_DIR, '御坊署_多発交差点一覧_R3-R7.xlsx')
wb.save(xlsx)
print('xlsx:', xlsx)

# --- csv (BOM付きUTF-8) ---
import csv
csvp = os.path.join(OUT_DIR, '御坊署_多発交差点一覧_R3-R7.csv')
with open(csvp, 'w', encoding='utf-8-sig', newline='') as f:
    w = csv.writer(f)
    w.writerow([h.replace('\n','') for h in headers])
    w.writerows(rows)
print('csv :', csvp)
print('交差点数:', len(rows))
